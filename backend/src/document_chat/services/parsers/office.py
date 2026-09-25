from __future__ import annotations

from pathlib import Path
from typing import List

from docx import Document
from docx.table import Table
from pptx import Presentation

from document_chat.services.parsers.chunks import make_chunk, split_text
from document_chat.services.parsers.kinds import TEXT_KIND, Chunk


def _table_text(rows: list[list[str]]) -> str:
    return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows if any(row))


def parse_docx(path: Path, record: dict) -> List[Chunk]:
    document = Document(path)
    sections: list[tuple[str, list[str]]] = [("", [])]
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            rows = [[cell.text for cell in row.cells] for row in block.rows]
            text = _table_text(rows)
            if text:
                sections[-1][1].append(text)
            continue
        text = block.text.strip()
        if not text:
            continue
        style = (block.style.name or "") if block.style is not None else ""
        if style.lower().startswith("heading"):
            sections.append((text, []))
        else:
            sections[-1][1].append(text)
    chunks: List[Chunk] = []
    for heading, lines in sections:
        body = "\n".join(lines).strip() or heading
        location = f"section '{heading}'" if heading else ""
        for piece in split_text(body):
            data = f"{heading}\n{piece}" if heading and piece != heading else piece
            chunks.append(make_chunk(record, TEXT_KIND, data, location=location))
    return chunks


def parse_pptx(path: Path, record: dict) -> List[Chunk]:
    presentation = Presentation(path)
    chunks: List[Chunk] = []
    for index, slide in enumerate(presentation.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                text = _table_text(rows)
            elif getattr(shape, "has_text_frame", False):
                text = shape.text_frame.text.strip()
            else:
                continue
            if text:
                lines.append(text)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"Speaker notes: {notes}")
        for piece in split_text("\n".join(lines)):
            chunks.append(make_chunk(record, TEXT_KIND, piece, location=f"slide {index}"))
    return chunks
